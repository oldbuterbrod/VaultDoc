import hashlib
import json
import mimetypes
import os
import random
import sys
import textwrap
import time
from datetime import datetime
from pathlib import Path

try:
    import requests
    from docx import Document
    from docx.shared import Pt
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.pdfgen import canvas
except Exception as exc:
    print(f"Не хватает зависимости: {exc}")
    print("Установи: python3 -m pip install requests reportlab python-docx")
    sys.exit(1)


TARGET_URL = os.environ.get("TARGET_URL", "http://127.0.0.1:8080").rstrip("/")
ADMIN_EMAIL = os.environ.get("ADMIN_EMAIL", "admin@vaultdoc.ru")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "AdminPass123!")
FOLDER_COUNT = int(os.environ.get("FOLDER_COUNT", "20"))
DOCUMENT_COUNT = int(os.environ.get("DOCUMENT_COUNT", "100"))

SCRIPT_DIR = Path(__file__).resolve().parent
STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")
OUT_DIR = SCRIPT_DIR / "generated_file_documents" / STAMP
OUT_DIR.mkdir(parents=True, exist_ok=True)

FOLDERS = [
    "01_Регламенты",
    "02_Политики_безопасности",
    "03_Инструкции",
    "04_Отчеты",
    "05_Проектная_документация",
    "06_Техническая_документация",
    "07_Договоры",
    "08_Акты",
    "09_Заявки",
    "10_Служебные_записки",
    "11_HR",
    "12_Финансы",
    "13_Разработка",
    "14_Инфраструктура",
    "15_Аудит",
    "16_Архив",
    "17_Внутренние_приказы",
    "18_Шаблоны",
    "19_Обучение",
    "20_Тестовые_материалы",
]

TITLE_TEMPLATES = [
    "Регламент обработки документов",
    "Политика разграничения доступа",
    "Инструкция пользователя VaultDoc",
    "Отчет по аудиту действий",
    "Описание архитектуры backend",
    "Описание архитектуры frontend",
    "Инструкция по загрузке файлов",
    "Заявка на предоставление доступа",
    "Акт проверки документа",
    "Служебная записка по безопасности",
    "Положение о хранении документов",
    "Памятка сотрудника",
    "Отчет о проверке прав доступа",
    "Описание сценариев тестирования",
    "Реестр внутренних документов",
    "Инструкция администратора",
    "План проверки отказоустойчивости",
    "Описание модели ролей",
    "Материалы обучения пользователей",
    "Тестовый документ файлового архива",
]

CATEGORY_PARAGRAPHS = {
    "Регламенты": [
        "Документ определяет порядок размещения, обновления и использования материалов в системе VaultDoc. Пользователь обязан указывать корректное название документа, выбирать соответствующую папку и загружать файл только разрешенного формата.",
        "В рамках регламента фиксируются основные этапы жизненного цикла документа: подготовка, загрузка, хранение, предоставление доступа, скачивание и удаление. Все значимые действия должны отражаться в журнале аудита.",
    ],
    "Политики безопасности": [
        "Документ описывает порядок разграничения доступа к материалам, размещенным в системе. Доступ к документам предоставляется на основании роли пользователя и объектных разрешений.",
        "Особое внимание уделяется контролю попыток несанкционированного доступа. Операции выдачи и отзыва прав фиксируются в журнале аудита для последующего анализа.",
    ],
    "Инструкции": [
        "Документ предназначен для пользователей системы VaultDoc и описывает типовой порядок работы с файловым хранилищем. Пользователь может просматривать доступные папки, открывать список документов и скачивать разрешенные файлы.",
        "Администратор может управлять пользователями, контролировать состояние учетных записей и анализировать журнал событий безопасности.",
    ],
    "Отчеты": [
        "Документ содержит демонстрационное описание результатов проверки работы системы. В ходе тестирования оцениваются доступность API, корректность авторизации и стабильность получения списков документов.",
        "Результаты используются для подтверждения работоспособности прототипа при наличии наполненного файлового хранилища.",
    ],
    "Проектная документация": [
        "Документ описывает проектные решения, принятые при разработке VaultDoc. Система построена как web-приложение с backend API, frontend-интерфейсом и reverse proxy.",
        "Для хранения данных используется PostgreSQL, а для управления структурой базы данных применяются миграции.",
    ],
    "Техническая документация": [
        "Документ содержит техническое описание компонентов системы. Backend реализует маршруты авторизации, пользователей, папок, документов, прав доступа и аудита.",
        "Frontend предоставляет интерфейс для работы с документами, журналом событий и административными разделами.",
    ],
    "Договоры": [
        "Документ используется как демонстрационный файл договорного типа. Он не содержит реальных обязательств и предназначен только для проверки загрузки и хранения файлов.",
        "В рамках тестового сценария документ размещается в файловом архиве и участвует в проверке отображения списка документов.",
    ],
    "Акты": [
        "Документ моделирует акт внутренней проверки. В нем фиксируется факт выполнения тестового действия и результат обработки документа системой.",
        "Использование таких документов позволяет проверить работу файлового хранилища на приближенных к реальным данных.",
    ],
    "Заявки": [
        "Документ моделирует заявку на выполнение операции в системе. В рамках тестового набора он используется для проверки отображения разных типов деловых документов.",
        "Заявка содержит сведения о цели обращения, ответственном подразделении и ожидаемом результате обработки.",
    ],
    "Служебные записки": [
        "Документ представляет собой служебную записку, подготовленную для внутреннего документооборота. Он используется для наполнения тестового файлового архива.",
        "Содержимое документа не является конфиденциальным и применяется только в рамках демонстрационного стенда.",
    ],
}

COMMON_PARAGRAPHS = [
    "Настоящий документ входит в демонстрационный набор данных VaultDoc. Набор используется для проверки загрузки, хранения, скачивания и отображения файловых документов в форматах PDF и DOCX.",
    "В рамках проверки моделируются типовые сценарии корпоративного документооборота: размещение документа в тематической папке, просмотр списка документов, получение метаданных, регистрация события загрузки и последующего скачивания в журнале аудита.",
    "Документ не содержит конфиденциальной информации и применяется только для функционального и нагрузочного тестирования прототипа.",
    "Система должна корректно обрабатывать документы разных категорий, сохранять сведения о файле, размер, контрольную сумму и дату загрузки.",
    "При работе с документом проверяется серверный контроль доступа. Пользователь может выполнять операции только в пределах своей роли и выданных объектных разрешений.",
    "Журнал аудита используется для последующего анализа действий пользователей и фиксации событий, связанных с документами, правами доступа и учетными записями.",
    "Демонстрационные документы помогают оценить поведение интерфейса при наличии наполненного хранилища и большого количества записей в списках.",
    "Полученные данные используются при подготовке результатов испытаний в выпускной квалификационной работе.",
]


def auth_headers(token):
    return {"Authorization": f"Bearer {token}"}


def login():
    response = requests.post(
        f"{TARGET_URL}/api/auth/login",
        data={"username": ADMIN_EMAIL, "password": ADMIN_PASSWORD},
        headers={"Content-Type": "application/x-www-form-urlencoded"},
        timeout=30,
        allow_redirects=False,
    )

    if response.status_code < 200 or response.status_code >= 300:
        print("Не удалось войти под admin")
        print(response.status_code)
        print(response.text)
        sys.exit(1)

    data = response.json()

    for key in ("access_token", "token", "accessToken"):
        value = data.get(key)
        if value:
            return value

    print("JWT не найден в ответе login")
    print(json.dumps(data, ensure_ascii=False, indent=2))
    sys.exit(1)


def get_openapi():
    for path in ("/openapi.json", "/api/openapi.json"):
        try:
            response = requests.get(f"{TARGET_URL}{path}", timeout=20, allow_redirects=False)
            if response.status_code == 200:
                return response.json()
        except Exception:
            pass
    return {}


def request_json(method, path, token, payload=None):
    url = f"{TARGET_URL}{path}"
    for attempt in range(5):
        response = requests.request(
            method,
            url,
            headers={**auth_headers(token), "Content-Type": "application/json"},
            json=payload,
            timeout=30,
            allow_redirects=False,
        )
        if response.status_code != 503:
            return response
        time.sleep(2 + attempt)
    return response


def existing_folders(token):
    result = []
    for path in ("/api/folders/", "/api/folders"):
        try:
            response = requests.get(f"{TARGET_URL}{path}", headers=auth_headers(token), timeout=30, allow_redirects=False)
            if response.status_code >= 200 and response.status_code < 300:
                data = response.json()
                if isinstance(data, list):
                    result.extend(data)
                elif isinstance(data, dict):
                    for key in ("items", "folders", "data", "results"):
                        value = data.get(key)
                        if isinstance(value, list):
                            result.extend(value)
        except Exception:
            pass
    return result


def create_folder(token, name):
    for folder in existing_folders(token):
        if folder.get("name") == name or folder.get("title") == name:
            return folder

    payloads = [
        {"name": name},
        {"name": name, "parent_id": None},
        {"title": name},
        {"title": name, "parent_id": None},
    ]

    for path in ("/api/folders/", "/api/folders"):
        for payload in payloads:
            response = request_json("POST", path, token, payload)
            if response.status_code >= 200 and response.status_code < 300:
                return response.json()
            if response.status_code in (400, 409):
                for folder in existing_folders(token):
                    if folder.get("name") == name or folder.get("title") == name:
                        return folder

    print(f"Не удалось создать папку: {name}")
    print(response.status_code)
    print(response.text)
    return None


def folder_identifier(folder):
    if not folder:
        return {}
    return {
        "folder_id": folder.get("id"),
        "folder_public_id": folder.get("public_id"),
        "parent_id": folder.get("id"),
        "parent_public_id": folder.get("public_id"),
    }


def category_from_folder(folder_name):
    value = folder_name.split("_", 1)[-1]
    value = value.replace("_", " ")
    mapping = {
        "Политики безопасности": "Политики безопасности",
        "Проектная документация": "Проектная документация",
        "Техническая документация": "Техническая документация",
        "Служебные записки": "Служебные записки",
    }
    return mapping.get(value, value)


def document_text(title, category, index, extension):
    paragraphs = []
    paragraphs.append(f"Демонстрационный документ VaultDoc")
    paragraphs.append(f"Наименование: {title}")
    paragraphs.append(f"Категория: {category}")
    paragraphs.append(f"Номер документа: VD-{index:03d}")
    paragraphs.append(f"Формат файла: {extension.upper()}")
    paragraphs.append(f"Версия: 1.0")
    paragraphs.append(f"Дата формирования: {datetime.now().strftime('%Y-%m-%d')}")
    paragraphs.extend(COMMON_PARAGRAPHS)

    category_text = CATEGORY_PARAGRAPHS.get(category, [])
    paragraphs.extend(category_text)

    while len(paragraphs) < 16:
        paragraphs.append(random.choice(COMMON_PARAGRAPHS))

    return paragraphs


def make_pdf(path, title, category, index):
    paragraphs = document_text(title, category, index, "pdf")
    c = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    x = 2 * cm
    y = height - 2 * cm

    c.setFont("Helvetica-Bold", 15)
    c.drawString(x, y, title[:80])
    y -= 1 * cm

    c.setFont("Helvetica", 10)

    for paragraph in paragraphs:
        lines = textwrap.wrap(paragraph, width=95)
        for line in lines:
            if y < 2 * cm:
                c.showPage()
                c.setFont("Helvetica", 10)
                y = height - 2 * cm
            c.drawString(x, y, line)
            y -= 0.45 * cm
        y -= 0.25 * cm

    c.save()


def make_docx(path, title, category, index):
    paragraphs = document_text(title, category, index, "docx")
    document = Document()
    document.add_heading(title, level=1)

    table = document.add_table(rows=5, cols=2)
    rows = [
        ("Категория", category),
        ("Номер документа", f"VD-{index:03d}"),
        ("Формат файла", "DOCX"),
        ("Версия", "1.0"),
        ("Дата формирования", datetime.now().strftime("%Y-%m-%d")),
    ]

    for row, values in zip(table.rows, rows):
        row.cells[0].text = values[0]
        row.cells[1].text = values[1]

    for paragraph in paragraphs:
        p = document.add_paragraph(paragraph)
        for run in p.runs:
            run.font.size = Pt(11)

    document.save(str(path))


def sha256(path):
    digest = hashlib.sha256()
    with open(path, "rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def openapi_upload_paths(spec):
    paths = []

    for path, methods in (spec.get("paths") or {}).items():
        if "documents" not in path:
            continue

        for method, operation in (methods or {}).items():
            if method.lower() != "post":
                continue

            request_body = operation.get("requestBody") or {}
            content = request_body.get("content") or {}

            if "multipart/form-data" in content:
                paths.append(path)

    fallback = [
        "/api/documents/upload",
        "/api/documents/",
        "/api/documents/{public_id}/upload",
        "/api/documents/{document_public_id}/upload",
        "/api/documents/{id}/upload",
        "/api/documents/{document_id}/upload",
    ]

    result = []

    for item in paths + fallback:
        if item not in result:
            result.append(item)

    return result


def openapi_metadata_paths(spec):
    paths = []

    for path, methods in (spec.get("paths") or {}).items():
        if "documents" not in path or "{" in path:
            continue

        for method, operation in (methods or {}).items():
            if method.lower() != "post":
                continue

            request_body = operation.get("requestBody") or {}
            content = request_body.get("content") or {}

            if "multipart/form-data" not in content:
                paths.append(path)

    fallback = [
        "/api/documents/",
    ]

    result = []

    for item in paths + fallback:
        if item not in result:
            result.append(item)

    return result


def format_path(path, document):
    public_id = document.get("public_id") or document.get("document_public_id") or document.get("id")
    document_id = document.get("id") or public_id

    replacements = {
        "{public_id}": str(public_id),
        "{document_public_id}": str(public_id),
        "{id}": str(document_id),
        "{document_id}": str(document_id),
    }

    result = path

    for key, value in replacements.items():
        result = result.replace(key, value)

    return result


def form_payload(title, folder):
    identifiers = folder_identifier(folder)
    payload = {
        "title": title,
        "name": title,
        "description": "Демонстрационный файловый документ для нагрузочного тестирования VaultDoc.",
        "folder_id": identifiers.get("folder_id") or "",
        "folder_public_id": identifiers.get("folder_public_id") or "",
        "parent_id": identifiers.get("parent_id") or "",
        "parent_public_id": identifiers.get("parent_public_id") or "",
    }

    return {key: str(value) for key, value in payload.items() if value is not None and value != ""}


def json_payload_variants(title, folder):
    identifiers = folder_identifier(folder)
    folder_id = identifiers.get("folder_id")
    folder_public_id = identifiers.get("folder_public_id")

    variants = [
        {"title": title, "name": title, "description": "Демонстрационный файловый документ для нагрузочного тестирования VaultDoc."},
        {"title": title, "description": "Демонстрационный файловый документ для нагрузочного тестирования VaultDoc."},
        {"name": title, "description": "Демонстрационный файловый документ для нагрузочного тестирования VaultDoc."},
    ]

    if folder_id:
        variants.extend([
            {"title": title, "folder_id": folder_id},
            {"name": title, "folder_id": folder_id},
            {"title": title, "name": title, "folder_id": folder_id},
        ])

    if folder_public_id:
        variants.extend([
            {"title": title, "folder_public_id": folder_public_id},
            {"name": title, "folder_public_id": folder_public_id},
            {"title": title, "name": title, "folder_public_id": folder_public_id},
        ])

    return variants


def upload_file_request(token, path, file_path, title, folder, file_field):
    mime = mimetypes.guess_type(str(file_path))[0] or "application/octet-stream"
    data = form_payload(title, folder)

    with open(file_path, "rb") as file:
        files = {
            file_field: (file_path.name, file, mime)
        }

        last_response = None
        for attempt in range(5):
            file.seek(0)
            last_response = requests.post(
                f"{TARGET_URL}{path}",
                headers=auth_headers(token),
                data=data,
                files=files,
                timeout=120,
                allow_redirects=False,
            )
            if last_response.status_code != 503:
                return last_response
            time.sleep(3 + attempt)
        return last_response


def direct_upload(token, upload_paths, file_path, title, folder):
    file_fields = ["file", "upload_file", "document_file", "uploaded_file"]

    errors = []

    for path in upload_paths:
        if "{" in path:
            continue

        for field in file_fields:
            response = upload_file_request(token, path, file_path, title, folder, field)

            if response.status_code >= 200 and response.status_code < 300:
                return True, response.json() if response.text else {}

            errors.append((path, field, response.status_code, response.text[:300]))

    return False, errors[-5:]


def create_document_metadata(token, metadata_paths, title, folder):
    errors = []

    for path in metadata_paths:
        for payload in json_payload_variants(title, folder):
            response = request_json("POST", path, token, payload)

            if response.status_code >= 200 and response.status_code < 300:
                try:
                    return response.json(), []
                except Exception:
                    return {}, []

            errors.append((path, response.status_code, response.text[:300]))

    return None, errors[-5:]


def upload_to_existing_document(token, upload_paths, document, file_path, title, folder):
    file_fields = ["file", "upload_file", "document_file", "uploaded_file"]
    errors = []

    path_candidates = []

    for path in upload_paths:
        if "{" in path:
            path_candidates.append(format_path(path, document))

    for value in (
        document.get("public_id"),
        document.get("document_public_id"),
        document.get("id"),
    ):
        if value:
            path_candidates.extend([
                f"/api/documents/{value}/upload",
                f"/api/documents/{value}/file",
                f"/api/documents/{value}/content",
            ])

    unique_paths = []

    for path in path_candidates:
        if path and path not in unique_paths and "{" not in path:
            unique_paths.append(path)

    for path in unique_paths:
        for field in file_fields:
            response = upload_file_request(token, path, file_path, title, folder, field)

            if response.status_code >= 200 and response.status_code < 300:
                return True, response.json() if response.text else {}

            errors.append((path, field, response.status_code, response.text[:300]))

    return False, errors[-5:]


def create_file_document(token, upload_paths, metadata_paths, file_path, title, folder):
    ok, response = direct_upload(token, upload_paths, file_path, title, folder)

    if ok:
        return True, response

    document, metadata_errors = create_document_metadata(token, metadata_paths, title, folder)

    if document:
        ok, response = upload_to_existing_document(token, upload_paths, document, file_path, title, folder)

        if ok:
            return True, response

        return False, {
            "stage": "upload_existing_document",
            "document": document,
            "errors": response,
        }

    return False, {
        "stage": "direct_or_metadata",
        "direct_errors": response,
        "metadata_errors": metadata_errors,
    }


def main():
    print(f"Цель: {TARGET_URL}")
    print("Логин под admin")

    token = login()
    spec = get_openapi()
    upload_paths = openapi_upload_paths(spec)
    metadata_paths = openapi_metadata_paths(spec)

    print("POST multipart candidates:")
    for path in upload_paths:
        print(f"  {path}")

    print("POST metadata candidates:")
    for path in metadata_paths:
        print(f"  {path}")

    print("Создание папок")

    selected_folders = FOLDERS[:FOLDER_COUNT]
    created_folders = []

    for name in selected_folders:
        folder = create_folder(token, name)
        if folder:
            created_folders.append(folder)
            print(f"  ok {name}")
        time.sleep(1)

    if not created_folders:
        print("Папки не созданы")
        sys.exit(1)

    print("Генерация и загрузка файлов")

    created = 0
    failed = 0
    report = {
        "target": TARGET_URL,
        "generated_at": STAMP,
        "folders": [],
        "documents": [],
        "failed": [],
    }

    for index in range(1, DOCUMENT_COUNT + 1):
        folder = created_folders[(index - 1) % len(created_folders)]
        folder_name = folder.get("name") or folder.get("title") or selected_folders[(index - 1) % len(selected_folders)]
        category = category_from_folder(folder_name)
        base_title = TITLE_TEMPLATES[(index - 1) % len(TITLE_TEMPLATES)]
        extension = "pdf" if index % 2 == 1 else "docx"
        title = f"{base_title} №{index:03d}"
        filename = f"{title}.{extension}".replace("/", "_")
        file_path = OUT_DIR / filename

        if extension == "pdf":
            make_pdf(file_path, title, category, index)
        else:
            make_docx(file_path, title, category, index)

        ok, response = create_file_document(token, upload_paths, metadata_paths, file_path, title, folder)

        item = {
            "index": index,
            "title": title,
            "filename": filename,
            "folder": folder_name,
            "format": extension,
            "size": file_path.stat().st_size,
            "sha256": sha256(file_path),
        }

        if ok:
            created += 1
            item["status"] = "created"
            item["response"] = response
            print(f"  ok {index:03d} {filename}")
        else:
            failed += 1
            item["status"] = "failed"
            item["error"] = response
            print(f"  fail {index:03d} {filename}")

        report["documents"].append(item)
        time.sleep(2)

    report["folders"] = created_folders
    report["created_documents"] = created
    report["failed_documents"] = failed

    report_path = OUT_DIR / "seed_file_documents_report.json"
    summary_path = OUT_DIR / "seed_file_documents_summary.txt"

    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    summary = [
        "Итоговая сводка наполнения файловыми документами",
        f"Цель: {TARGET_URL}",
        f"Папка генерации: {OUT_DIR}",
        f"Создано папок: {len(created_folders)}",
        f"Создано файловых документов: {created}",
        f"Ошибок загрузки: {failed}",
        f"PDF: {sum(1 for item in report['documents'] if item['format'] == 'pdf' and item['status'] == 'created')}",
        f"DOCX: {sum(1 for item in report['documents'] if item['format'] == 'docx' and item['status'] == 'created')}",
    ]

    summary_path.write_text("\n".join(summary), encoding="utf-8")

    print()
    print("\n".join(summary))
    print()
    print(summary_path)

    if failed:
        sys.exit(2)


if __name__ == "__main__":
    main()
